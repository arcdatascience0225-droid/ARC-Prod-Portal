from fastapi import APIRouter, Depends
from sqlalchemy.orm import Session
from app.db.database import get_db
from app.core.deps import get_current_user
from app.services.ai_assistant_service import AIAssistantService
from app.schemas import ai_assistant as schemas

router = APIRouter(prefix="/ai", tags=["AI Assistant"])


@router.post("/chat", response_model=schemas.ChatResponse)
def chat(payload: schemas.ChatRequest, db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    response, tokens = AIAssistantService(db).chat(current_user.id, payload.message, payload.sessionContext)
    return schemas.ChatResponse(response=response, tokensUsed=tokens)


@router.post("/resume/generate")
def generate_resume(payload: schemas.ResumeGenerateRequest, db: Session = Depends(get_db),
                     current_user=Depends(get_current_user)):
    resume = AIAssistantService(db).generate_resume(current_user.id, payload.profile, payload.achievements, payload.targetRole)
    return {"id": resume.id, "content": resume.content, "aiGeneratedText": resume.ai_generated_text, "version": resume.version}


@router.post("/resume/improve")
def improve_resume(payload: schemas.ResumeImproveRequest, db: Session = Depends(get_db),
                    current_user=Depends(get_current_user)):
    text, tokens = AIAssistantService(db).improve_resume(current_user.id, payload.resumeText, payload.targetRole)
    return {"suggestions": text, "tokensUsed": tokens}


@router.post("/resume/export/pdf", summary="Export resume text as a downloadable PDF")
def export_resume_pdf(payload: dict, current_user=Depends(get_current_user)):
    from fastapi.responses import StreamingResponse
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    import io

    resume_text = payload.get("resumeText", "")
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    body_style = styles["BodyText"]
    body_style.leading = 14

    story = []
    for line in resume_text.split("\n"):
        if not line.strip():
            story.append(Spacer(1, 8))
            continue
        # Escape characters reportlab's mini-markup would otherwise choke on.
        safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe_line, body_style))

    doc.build(story)
    buffer.seek(0)
    return StreamingResponse(
        buffer, media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=resume.pdf"},
    )


@router.post("/resume/export/docx", summary="Export resume text as a downloadable Word document")
def export_resume_docx(payload: dict, current_user=Depends(get_current_user)):
    from fastapi.responses import StreamingResponse
    from docx import Document
    import io

    resume_text = payload.get("resumeText", "")
    doc = Document()
    for line in resume_text.split("\n"):
        if not line.strip():
            doc.add_paragraph("")
            continue
        doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=resume.docx"},
    )


@router.post("/career-guidance")
def career_guidance(payload: schemas.CareerGuidanceRequest, db: Session = Depends(get_db),
                     current_user=Depends(get_current_user)):
    text, tokens = AIAssistantService(db).career_guidance(current_user.id, payload.question, payload.interestArea)
    return {"response": text, "tokensUsed": tokens}


@router.post("/study-plan")
def generate_study_plan(payload: schemas.StudyPlanRequest, db: Session = Depends(get_db),
                         current_user=Depends(get_current_user)):
    plan = AIAssistantService(db).generate_study_plan(current_user.id, payload.goal, payload.currentLevel, payload.hoursPerWeek)
    return {"id": plan.id, "goal": plan.goal, "planData": plan.plan_data}


@router.get("/strength-weakness", response_model=schemas.StrengthWeaknessOut)
def strength_weakness(db: Session = Depends(get_db), current_user=Depends(get_current_user)):
    return AIAssistantService(db).strength_weakness_analysis(current_user.id)
